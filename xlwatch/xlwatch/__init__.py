import os
import math
import time
import json
import socket
import openpyxl
import subprocess
import watchdog.events
import watchdog.observers
import threading

def get_table_with_key(table: dict, key: str):
    key_split = key.split('.')

    cur = table
    for name in key_split:
        if not cur.get(name):
            return None
        cur = cur[name]
    return cur

def set_table_with_key(table: dict, key: str, value):
    key_split = key.split('.')

    cur = table
    for name in key_split[:-1]:
        if not cur.get(name):
            cur[name] = {}
        cur = cur[name]
    cur[key_split[-1]] = value

def stringify_as_lua_type(table, indent = 0):
    tmp = '{\n'
    for name in table:
        if type(table[name]) is dict:
            tmp += ('\t' * (indent + 1)) + f'{name}: {stringify_as_lua_type(table[name], indent + 1)},\n'
        else:
            tmp += ('\t' * (indent + 1)) + f'{name}: {table[name]},\n'
    
    if len(table) >= 1:
        tmp = tmp[:-2]
    tmp += '\n' + ('\t' * indent) + '}'

    return tmp

def stringify_as_lua_table(table, indent = 0):
    tmp = '{\n'
    for name in table:
        if type(table[name]) is dict:
            tmp += ('\t' * (indent + 1)) + f'{name} = {stringify_as_lua_table(table[name], indent + 1)},\n'
        elif table[name] is None:
            tmp += ('\t' * (indent + 1)) + f'{name} = nil,\n'
        else:
            tmp += ('\t' * (indent + 1)) + f'{name} = {json.dumps(table[name], ensure_ascii=False)},\n'
    
    if len(table) >= 1:
        tmp = tmp[:-2]
    tmp += '\n' + ('\t' * indent) + '}'
    
    return tmp

def to_number_big(x: str):
    if type(x) is not str:
        return x
    
    units = ["", "K", "M", "B", "T", "Qa", "Qi", "Sx", "Sp", "Oc"]

    for i in range(1, len(units)):
        if x[-len(units[i]):].lower() == units[i].lower():
            return float(x[:-len(units[i])]) * math.pow(1000, i)

    return float(x)

def is_port_in_use(port: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        return s.connect_ex(('localhost', port)) == 0

def run_rojo():
    subprocess.run("xlwatch/rojo.exe serve", cwd="../")
    pass

class XlWatchEventHandler(watchdog.events.PatternMatchingEventHandler):
    def __init__(self, xl_path : str, output_path : str):
        watchdog.events.PatternMatchingEventHandler.__init__(self, patterns=['*.xlsx'], ignore_directories=True, case_sensitive=False)
        self.output_path = output_path
        self.type_table = {}
        self.data_table = {}
        self.key_type = {}

        for file in os.scandir(xl_path):
            if file.is_file() and file.path.endswith('.xlsx') and not '~$' in file.path:
                self.read(file.path)
        self.write()
        print(f'OK')

    def read(self, path : str):
        try:
            workbook = openpyxl.load_workbook(path, data_only=True)
            for sheetname in workbook.sheetnames:
                self.type_table[sheetname] = {}
                self.data_table[sheetname] = {}

                sheet = workbook.get_sheet_by_name(sheetname)

                self.key_type[sheetname] = sheet.cell(2, 1).value
                
                for i in range(2, sheet.max_column + 1):
                    if sheet.cell(1, i).value == None:
                        continue
                    
                    _type = sheet.cell(2, i).value
                    if _type == 'int':
                        set_table_with_key(self.type_table[sheetname], sheet.cell(1, i).value, 'number')
                    elif _type == 'float':
                        set_table_with_key(self.type_table[sheetname], sheet.cell(1, i).value, 'number')
                    elif _type == 'string':
                        set_table_with_key(self.type_table[sheetname], sheet.cell(1, i).value, 'string')
                    elif _type == 'bool':
                        set_table_with_key(self.type_table[sheetname], sheet.cell(1, i).value, 'boolean')

                for i in range(3, sheet.max_row + 1):
                    id = sheet.cell(i, 1).value
                    if id == None:
                        continue
                    if type(id) is int:
                        id = f"[{id}]"
                    self.data_table[sheetname][id] = {}

                    for j in range(2, sheet.max_column + 1):
                        if sheet.cell(1, j).value == None:
                            continue

                        _type = sheet.cell(2, j).value
                        if sheet.cell(i, j).value == None:
                            set_table_with_key(self.data_table[sheetname][id], sheet.cell(1, j).value, None)
                        elif _type == 'int':
                            set_table_with_key(self.data_table[sheetname][id], sheet.cell(1, j).value, int(to_number_big(sheet.cell(i, j).value)))
                        elif _type == 'float':
                            set_table_with_key(self.data_table[sheetname][id], sheet.cell(1, j).value, to_number_big(sheet.cell(i, j).value))
                        elif _type == 'string':
                            set_table_with_key(self.data_table[sheetname][id], sheet.cell(1, j).value, str(sheet.cell(i, j).value))
                        elif _type == 'bool':
                            set_table_with_key(self.data_table[sheetname][id], sheet.cell(1, j).value, True if sheet.cell(i, j).value == True else False)

        except PermissionError:
            pass

    def write(self):
        with open(self.output_path, 'wb') as f:
            tmp = ''

            for _type in self.type_table:
                tmp += f'export type {_type}_Type = {stringify_as_lua_type(self.type_table[_type])}\n\n'
            
            tmp += f'export type Table_Type = {{\n'
            for _type in self.type_table:
                tmp += f'\t{_type}: {{[{self.key_type[_type]}]: {_type}_Type}},\n'
            if len(self.type_table) >= 1:
                tmp = tmp[:-1]
            tmp += '\n}\n\n'

            tmp += f'local Table : Table_Type = {stringify_as_lua_table(self.data_table)}\n\n'

            tmp += f'return Table'

            f.write(tmp.encode(encoding='utf-8'))

    def on_modified(self, event : watchdog.events.FileModifiedEvent):
        if not '~$' in event.src_path:
            self.read(event.src_path)
            print(f'{event.src_path} 적용 중...', end = '')
            self.write()
            print(f'OK')

class XlWatchObserver:
    def __init__(self, xl_path : str, output_path : str):
        if not is_port_in_use(34872):
            threading.Thread(target=run_rojo, daemon=True).start()

        event_handler = XlWatchEventHandler(xl_path, output_path)
        observer = watchdog.observers.Observer()
        observer.schedule(event_handler, './')
        observer.start()

        try:
            while True:
                time.sleep(1)
        except KeyboardInterrupt:
            observer.stop()

        observer.join()
